"""Sorteia N avaliações e gera a planilha de rotulagem manual (.xlsx).

Gera dois arquivos:
  - amostra_validacao.csv : as mesmas avaliações (com ids) para a IA classificar;
  - planilha_rotulagem.xlsx : modelo para cada rotulador preencher SEM ver a IA.

Exemplo:
  python scripts/gerar_planilha_rotulagem.py data/olist_500.csv --n 100
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from feedback_radar.anonimizacao import anonimizar  # noqa: E402
from feedback_radar.leitura_csv import ler_csv  # noqa: E402
from feedback_radar.taxonomia import TEMAS, TEMAS_VALIDOS  # noqa: E402

INSTRUCOES = [
    "MANUAL DE ROTULAGEM — Feedback Radar",
    "",
    "Objetivo: criar um 'gabarito' humano para medir se a IA classifica bem as avaliações.",
    "",
    "REGRAS PARA EVITAR VIÉS",
    "1. NÃO rode a IA nessas avaliações antes de terminar a rotulagem, e não olhe o resultado dela.",
    "2. Cada avaliação deve ser rotulada por DUAS pessoas, de forma independente, em cópias separadas "
    "(ex.: rotulador_A.xlsx e rotulador_B.xlsx). Não converse sobre os casos durante a rotulagem.",
    "3. Leia SÓ o texto. A nota (estrelas) foi omitida de propósito: queremos avaliar o texto.",
    "4. Siga as definições da aba 'temas'. Na dúvida, marque 'duvida' = sim e explique em 'observacoes'.",
    "5. Rotule na ordem da planilha (ela já foi embaralhada) e faça pausas a cada ~30 avaliações.",
    "6. Depois, rode scripts/avaliar_validacao.py com --rotulador-b para medir a concordância entre "
    "vocês (Kappa). Em reunião, resolvam as divergências e salvem o consenso como gabarito.xlsx.",
    "",
    "COMO PREENCHER",
    "- sentimento: positivo, neutro ou negativo (sentimento GERAL do cliente). Mista equilibrada = neutro.",
    "- tema_*: marque 'x' em TODOS os temas mencionados (pelo menos um). 'outros' só se nenhum servir.",
    "- intensidade: 1 (indiferente) a 5 (emoção extrema, ex.: 'nunca mais volto').",
    "- Ironia ('Parabéns, esperei 2 horas'): rotule pelo sentido REAL (negativo).",
    "",
    "PRIVACIDADE: os textos já foram anonimizados. Não compartilhe esta planilha fora do grupo.",
]


def main():
    p = argparse.ArgumentParser()
    p.add_argument("entrada", help="CSV de avaliações (coluna 'texto')")
    p.add_argument("--n", type=int, default=100)
    p.add_argument("--semente", type=int, default=2026)
    p.add_argument("--pasta", default="data/rotulagem")
    args = p.parse_args()

    df, _ = ler_csv(args.entrada)
    amostra = df.sample(n=min(args.n, len(df)), random_state=args.semente).reset_index(drop=True)
    amostra["id"] = range(1, len(amostra) + 1)  # ids novos, de 1 a N
    amostra["texto"] = amostra["texto"].map(anonimizar)

    pasta = Path(args.pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    colunas = ["id", "texto"] + [c for c in ("data", "nota", "fonte") if c in amostra]
    amostra[colunas].to_csv(pasta / "amostra_validacao.csv", index=False, encoding="utf-8-sig")

    wb = Workbook()
    ws_inst = wb.active
    ws_inst.title = "instrucoes"
    for linha in INSTRUCOES:
        ws_inst.append([linha])
    ws_inst["A1"].font = Font(bold=True, size=14)
    ws_inst.column_dimensions["A"].width = 130

    ws_temas = wb.create_sheet("temas")
    ws_temas.append(["coluna", "tema", "definição"])
    for chave, (nome, definicao) in TEMAS.items():
        ws_temas.append([f"tema_{chave}", nome, definicao])
    ws_temas.column_dimensions["A"].width = 26
    ws_temas.column_dimensions["C"].width = 110

    ws = wb.create_sheet("rotulagem")
    cabecalho = (["id", "texto", "sentimento"] + [f"tema_{t}" for t in TEMAS_VALIDOS]
                 + ["intensidade", "duvida", "observacoes"])
    ws.append(cabecalho)
    for r in amostra.itertuples():
        ws.append([r.id, r.texto])

    destaque = PatternFill("solid", fgColor="DDEBF7")
    for celula in ws[1]:
        celula.font = Font(bold=True)
        celula.fill = destaque
    ws.freeze_panes = "C2"
    ws.column_dimensions["B"].width = 80
    for linha in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        linha[0].alignment = Alignment(wrap_text=True, vertical="top")

    # Listas suspensas reduzem erros de digitação
    n = len(amostra) + 1
    validacoes = {
        "C": '"positivo,neutro,negativo"',
        "intensidade": '"1,2,3,4,5"',
        "duvida": '"sim,não"',
    }
    dv = DataValidation(type="list", formula1=validacoes["C"], allow_blank=True)
    dv.add(f"C2:C{n}")
    ws.add_data_validation(dv)
    for i, nome in enumerate(cabecalho, start=1):
        letra = ws.cell(row=1, column=i).column_letter
        if nome.startswith("tema_"):
            dv_tema = DataValidation(type="list", formula1='"x"', allow_blank=True)
            dv_tema.add(f"{letra}2:{letra}{n}")
            ws.add_data_validation(dv_tema)
            ws.column_dimensions[letra].width = 12
        elif nome in ("intensidade", "duvida"):
            dv_extra = DataValidation(type="list", formula1=validacoes[nome], allow_blank=True)
            dv_extra.add(f"{letra}2:{letra}{n}")
            ws.add_data_validation(dv_extra)

    wb.active = wb.index(ws_inst)
    wb.save(pasta / "planilha_rotulagem.xlsx")
    print(f"{len(amostra)} avaliações sorteadas.\n"
          f"  -> {pasta / 'amostra_validacao.csv'} (para a IA classificar DEPOIS da rotulagem)\n"
          f"  -> {pasta / 'planilha_rotulagem.xlsx'} (copie uma por rotulador)")


if __name__ == "__main__":
    main()
