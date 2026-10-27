"""Divide a planilha de rotulagem entre as duplas: um arquivo pronto para cada pessoa.

Com 6 integrantes: 3 duplas, cada uma com ~33 avaliações. As duas pessoas de uma dupla
recebem as MESMAS avaliações e rotulam separadas (é o que permite medir a concordância).

  python scripts/dividir_rotulagem.py --duplas 3

Gera data/rotulagem/para_preencher/dupla1_pessoaA.xlsx, dupla1_pessoaB.xlsx, ...
"""
import argparse
from pathlib import Path

from openpyxl import load_workbook


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--planilha", default="data/rotulagem/planilha_rotulagem.xlsx")
    p.add_argument("--duplas", type=int, default=3)
    p.add_argument("--pasta", default="data/rotulagem/para_preencher")
    args = p.parse_args()

    total = load_workbook(args.planilha)["rotulagem"].max_row - 1
    tamanho = -(-total // args.duplas)  # divisão arredondando para cima
    pasta = Path(args.pasta)
    pasta.mkdir(parents=True, exist_ok=True)

    for dupla in range(args.duplas):
        inicio, fim = dupla * tamanho + 1, min((dupla + 1) * tamanho, total)  # posições das linhas
        for pessoa in ("A", "B"):
            wb = load_workbook(args.planilha)  # cópia nova a cada arquivo
            ws = wb["rotulagem"]
            # Apaga de baixo para cima as linhas que não são desta dupla (linha 1 = cabeçalho)
            for linha in range(ws.max_row, 1, -1):
                if not inicio <= linha - 1 <= fim:
                    ws.delete_rows(linha)
            wb["instrucoes"].append([])
            wb["instrucoes"].append([f"Este arquivo: dupla {dupla + 1}, pessoa {pessoa}. "
                                     f"Avaliações {inicio} a {fim}. Rotule sozinho(a), sem conversar "
                                     "com a outra pessoa da dupla até as duas terminarem."])
            destino = pasta / f"dupla{dupla + 1}_pessoa{pessoa}.xlsx"
            wb.save(destino)
            print(f"{destino}: avaliações {inicio} a {fim}")


if __name__ == "__main__":
    main()
