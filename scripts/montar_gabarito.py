"""Junta as planilhas preenchidas pelas duplas e monta o gabarito.

Para cada dupla, compara a pessoa A com a pessoa B:
- onde as duas concordam no sentimento, o rótulo vai direto para o gabarito
  (temas: os marcados pelas duas; se não houver nenhum em comum, entra como divergência);
- onde discordam, a linha fica em branco no gabarito e aparece em divergencias.xlsx
  com os dois rótulos, para a dupla decidir junto na reunião.

  python scripts/montar_gabarito.py

Depois da reunião, preencham as linhas em branco do gabarito.xlsx (aba "rotulagem")
e rodem avaliar_validacao.py. Linhas que continuarem em branco ficam fora das métricas.
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sklearn.metrics import cohen_kappa_score  # noqa: E402

from feedback_radar.classificador import lista_temas  # noqa: E402
from feedback_radar.taxonomia import TEMAS_VALIDOS  # noqa: E402
from feedback_radar.validacao import carregar_gabarito  # noqa: E402


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--pasta", default="data/rotulagem/para_preencher")
    p.add_argument("--modelo", default="data/rotulagem/planilha_rotulagem.xlsx",
                   help="planilha vazia usada como base do gabarito")
    p.add_argument("--saida", default="data/rotulagem")
    args = p.parse_args()

    pasta = Path(args.pasta)
    decisoes, divergencias, resumo = {}, [], []
    for arquivo_a in sorted(pasta.glob("dupla*_pessoaA.xlsx")):
        arquivo_b = arquivo_a.with_name(arquivo_a.name.replace("pessoaA", "pessoaB"))
        if not arquivo_b.exists():
            sys.exit(f"Falta {arquivo_b.name}.")
        a, b = carregar_gabarito(arquivo_a), carregar_gabarito(arquivo_b)
        juntos = a.merge(b, on=["id", "texto"], suffixes=("_a", "_b"))
        if juntos.empty:
            print(f"{arquivo_a.stem[:6]}: nenhuma avaliação preenchida pelas duas pessoas ainda.")
            continue

        for r in juntos.itertuples():
            temas_a, temas_b = set(lista_temas(r.temas_humano_a)), set(lista_temas(r.temas_humano_b))
            comuns = temas_a & temas_b
            if r.sentimento_humano_a == r.sentimento_humano_b and comuns:
                decisoes[r.id] = (r.sentimento_humano_a, comuns)
            else:
                divergencias.append({
                    "dupla": arquivo_a.stem[:6], "id": r.id, "texto": r.texto,
                    "sentimento_A": r.sentimento_humano_a, "sentimento_B": r.sentimento_humano_b,
                    "temas_A": ", ".join(sorted(temas_a)), "temas_B": ", ".join(sorted(temas_b)),
                    "decisao_da_dupla": "",
                })
        kappa = cohen_kappa_score(juntos["sentimento_humano_a"], juntos["sentimento_humano_b"])
        resumo.append(f"{arquivo_a.stem[:6]}: {len(juntos)} avaliações, concordância de sentimento "
                      f"{(juntos['sentimento_humano_a'] == juntos['sentimento_humano_b']).mean():.0%}, "
                      f"Kappa {kappa:.2f}")

    if not resumo:
        sys.exit(f"Nenhuma planilha preenchida encontrada em {pasta}.")

    # Gabarito no mesmo formato da planilha de rotulagem (é o que carregar_gabarito lê)
    wb = load_workbook(args.modelo)
    ws = wb["rotulagem"]
    cabecalho = [c.value for c in ws[1]]
    col_sent = cabecalho.index("sentimento") + 1
    for linha in range(2, ws.max_row + 1):
        id_ = ws.cell(linha, 1).value
        if id_ in decisoes:
            sentimento, temas = decisoes[id_]
            ws.cell(linha, col_sent).value = sentimento
            for tema in TEMAS_VALIDOS:
                ws.cell(linha, cabecalho.index(f"tema_{tema}") + 1).value = "x" if tema in temas else None
    saida = Path(args.saida)
    wb.save(saida / "gabarito.xlsx")
    pd.DataFrame(divergencias).to_excel(saida / "divergencias.xlsx", index=False)

    print("\n".join(resumo))
    print(f"\n{len(decisoes)} avaliações decididas automaticamente (as duas pessoas concordaram).")
    print(f"{len(divergencias)} para decidir na reunião: {saida / 'divergencias.xlsx'}")
    print(f"Gabarito salvo em {saida / 'gabarito.xlsx'}; preencham nele as linhas em branco.")


if __name__ == "__main__":
    main()
