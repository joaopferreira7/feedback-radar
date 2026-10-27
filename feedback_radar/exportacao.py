"""Arquivos para download pensados para o Excel em português.

O Excel brasileiro usa vírgula como separador decimal e, por isso, espera ponto e
vírgula entre as colunas de um CSV. Um CSV separado por vírgulas abre com tudo
numa coluna só. Aqui geramos:
- um .xlsx formatado (o que o usuário deve abrir no Excel);
- um CSV com ";" e decimal "," que o Excel abre certo e o app aceita de volta.
"""
import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .taxonomia import NOMES_TEMAS

# Ordem e nomes das colunas na planilha
COLUNAS = {
    "id": "ID",
    "texto": "Avaliação",
    "sentimento": "Sentimento",
    "temas": "Temas",
    "intensidade": "Intensidade (1 a 5)",
    "resumo_curto": "Resumo",
    "nota": "Nota",
    "data": "Data",
    "fonte": "Fonte",
    "modelo": "Modelo",
    "status": "Status",
}
LARGURAS = {"Avaliação": 70, "Resumo": 45, "Temas": 30, "Sentimento": 13, "Data": 12,
            "Intensidade (1 a 5)": 11, "Nota": 7, "ID": 6, "Fonte": 12, "Modelo": 28, "Status": 9}
# Tons claros para não brigar com o texto (mesma família de cores do app)
FUNDO_SENTIMENTO = {"positivo": "DDE8DF", "neutro": "ECE9E3", "negativo": "F3DCD8"}


def _temas_por_extenso(valor) -> str:
    if not isinstance(valor, str) or not valor:
        return ""
    return ", ".join(NOMES_TEMAS.get(t, t) for t in valor.split(";") if t)


def _inteiros(df: pd.DataFrame) -> pd.DataFrame:
    # Colunas com algum vazio viram float (4.0); para o Excel queremos 4
    for col in ("id", "intensidade"):
        if col in df:
            df[col] = pd.to_numeric(df[col], errors="coerce").astype("Int64")
    if "nota" in df:  # nota 4.5 continua decimal; 5.0 vira 5
        nota = pd.to_numeric(df["nota"], errors="coerce")
        if (nota.dropna() % 1 == 0).all():
            df["nota"] = nota.astype("Int64")
    return df


def para_excel(df: pd.DataFrame, nome_aba: str = "Avaliações") -> bytes:
    """Planilha .xlsx pronta para ler: cabeçalho fixo, filtros, larguras e cores por sentimento."""
    tabela = _inteiros(df.copy())
    # Primeiro as colunas conhecidas, na ordem acima; depois as demais (ex.: as da tabela de erros).
    # texto_enviado é a versão anonimizada do texto e só duplicaria a coluna.
    colunas = [c for c in COLUNAS if c in tabela] + \
              [c for c in tabela if c not in COLUNAS and c != "texto_enviado"]
    if "status" in tabela and (tabela["status"] == "ok").all():
        colunas.remove("status")  # só interessa quando alguma avaliação falhou
    tabela = tabela[colunas]
    for col in ("temas", "temas_humano", "temas_ia"):
        if col in tabela:
            tabela[col] = tabela[col].map(_temas_por_extenso)
    if "data" in tabela:
        tabela["data"] = pd.to_datetime(tabela["data"], errors="coerce").dt.date
    tabela = tabela.rename(columns=COLUNAS)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as escritor:
        tabela.to_excel(escritor, index=False, sheet_name=nome_aba)
        ws = escritor.sheets[nome_aba]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for celula in ws[1]:
            celula.font = Font(bold=True, color="F4F1EA")
            celula.fill = PatternFill("solid", fgColor="2F4858")
            celula.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 30

        nomes = [c.value for c in ws[1]]
        for i, nome in enumerate(nomes, start=1):
            ws.column_dimensions[get_column_letter(i)].width = LARGURAS.get(nome, 15)
        col_sent = nomes.index("Sentimento") + 1 if "Sentimento" in nomes else None
        col_data = nomes.index("Data") + 1 if "Data" in nomes else None

        for linha in ws.iter_rows(min_row=2):
            for celula in linha:
                celula.alignment = Alignment(vertical="top", wrap_text=True)
            if col_data:
                linha[col_data - 1].number_format = "DD/MM/YYYY"
            if col_sent:
                cor = FUNDO_SENTIMENTO.get(str(linha[col_sent - 1].value))
                if cor:
                    linha[col_sent - 1].fill = PatternFill("solid", fgColor=cor)
    return buffer.getvalue()


def para_csv(df: pd.DataFrame) -> bytes:
    """CSV no padrão do Excel brasileiro (";" entre colunas, "," nos decimais, datas dd/mm/aaaa).

    Mantém os nomes internos das colunas para o arquivo poder ser enviado de volta ao app.
    """
    tabela = _inteiros(df.copy())
    if "data" in tabela:
        tabela["data"] = pd.to_datetime(tabela["data"], errors="coerce").dt.strftime("%d/%m/%Y")
    return tabela.to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")
